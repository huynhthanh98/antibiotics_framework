#!/usr/bin/env python3
"""
Xuất Excel so sánh Task1(cells) vs Task2(patches) cho Image-Level (Soft Vote),
lọc theo danh sách backbone chọn trước.

Input: reports/task1_vs_task2_resistant_metrics.csv
Output: reports/softvote_cells_vs_patches_selected_backbones.xlsx
"""

from __future__ import annotations

import csv
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
IN_CSV = ROOT / "reports" / "task1_vs_task2_resistant_metrics.csv"
OUT_XLSX = ROOT / "reports" / "softvote_cells_vs_patches_selected_backbones.xlsx"

LEVEL = "Image-Level (Soft Vote)"

# 15 backbone có Cells > Patches cho Recall(R) trên cả 3 thuốc,
# theo thống kê trước đó, nhưng loại efficientnet_b5 theo yêu cầu.
SELECTED_BACKBONES = [
    "convnext_base",
    "convnext_small",
    "convnext_tiny",
    "convnextv2_base",
    "convnextv2_tiny",
    "densenet121",
    "densenet201",
    "resnet101",
    "resnet50",
    "swin_base",
    "swin_small",
    "swin_tiny",
    "vgg16",
    "vgg19",
]

METRICS = [
    ("recall_r", "Recall(R)"),
    ("prec_r", "Precision(R)"),
    ("f1_r", "F1(R)"),
    ("acc", "Accuracy"),
    ("auc", "AUC"),
]


def main() -> None:
    if not IN_CSV.exists():
        raise SystemExit(f"Không thấy input CSV: {IN_CSV}")

    rows = []
    with IN_CSV.open(newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if r["level"] != LEVEL:
                continue
            if r["backbone"] not in set(SELECTED_BACKBONES):
                continue
            rows.append(r)

    # Index: (drug, backbone) -> row
    idx: dict[tuple[str, str], dict] = {(r["drug"], r["backbone"]): r for r in rows}

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    fill_t1_win = PatternFill("solid", fgColor="DBEAFE")  # xanh dương
    fill_t2_win = PatternFill("solid", fgColor="DCFCE7")  # xanh lá
    fill_tie = PatternFill("solid", fgColor="FEF9C3")  # vàng
    fill_lose = PatternFill("solid", fgColor="F1F5F9")  # xám nhạt
    font_win = Font(bold=True)
    font_lose = Font(color="64748B")
    font_hdr = Font(bold=True)
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def apply_pair(ws, row_i: int, col_t1: int, col_t2: int, a: float, b: float) -> None:
        c1 = ws.cell(row=row_i, column=col_t1)
        c2 = ws.cell(row=row_i, column=col_t2)
        if abs(a - b) < 1e-12:
            for c in (c1, c2):
                c.fill = fill_tie
                c.font = Font(bold=True)
        elif a > b:
            c1.fill, c1.font = fill_t1_win, font_win
            c2.fill, c2.font = fill_lose, font_lose
        else:
            c2.fill, c2.font = fill_t2_win, font_win
            c1.fill, c1.font = fill_lose, font_lose

    def make_sheet(wb: Workbook, drug: str) -> None:
        ws = wb.create_sheet(title=drug)
        # Header: backbone + (metric T1, metric T2) lặp
        heads = ["Backbone"]
        for _, label in METRICS:
            heads.extend([f"{label} (Cells)", f"{label} (Patches)"])

        for j, h in enumerate(heads, start=1):
            cell = ws.cell(row=1, column=j, value=h)
            cell.font = font_hdr
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for i, backbone in enumerate(SELECTED_BACKBONES, start=2):
            ws.cell(row=i, column=1, value=backbone).border = border
            ws.cell(row=i, column=1).alignment = Alignment(horizontal="left")
            r = idx.get((drug, backbone))
            if not r:
                # để trống nếu thiếu log
                for j in range(2, len(heads) + 1):
                    ws.cell(row=i, column=j, value=None).border = border
                continue

            col = 2
            for key, _ in METRICS:
                a = float(r[f"t1_{key}"])
                b = float(r[f"t2_{key}"])
                c1 = ws.cell(row=i, column=col, value=round(a, 6))
                c2 = ws.cell(row=i, column=col + 1, value=round(b, 6))
                for c in (c1, c2):
                    c.number_format = "0.0000"
                    c.border = border
                    c.alignment = Alignment(horizontal="right")
                apply_pair(ws, i, col, col + 1, a, b)
                col += 2

        # widths + freeze
        ws.column_dimensions["A"].width = 18
        for j in range(2, len(heads) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 16
        ws.freeze_panes = "A2"

    wb = Workbook()
    # Xóa default sheet
    wb.remove(wb.active)

    # Sheet tổng hợp (3 thuốc x 14 backbone)
    ws_all = wb.create_sheet(title="All")
    heads = ["Drug", "Backbone"]
    for _, label in METRICS:
        heads.extend([f"{label} (Cells)", f"{label} (Patches)"])
    for j, h in enumerate(heads, start=1):
        c = ws_all.cell(row=1, column=j, value=h)
        c.font = font_hdr
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    r_i = 2
    for drug in ["AZI", "CIP", "TE"]:
        for backbone in SELECTED_BACKBONES:
            ws_all.cell(row=r_i, column=1, value=drug).border = border
            ws_all.cell(row=r_i, column=2, value=backbone).border = border
            ws_all.cell(row=r_i, column=1).alignment = Alignment(horizontal="left")
            ws_all.cell(row=r_i, column=2).alignment = Alignment(horizontal="left")
            r = idx.get((drug, backbone))
            col = 3
            for key, _ in METRICS:
                if r:
                    a = float(r[f"t1_{key}"])
                    b = float(r[f"t2_{key}"])
                    c1 = ws_all.cell(row=r_i, column=col, value=round(a, 6))
                    c2 = ws_all.cell(row=r_i, column=col + 1, value=round(b, 6))
                    for c in (c1, c2):
                        c.number_format = "0.0000"
                        c.border = border
                        c.alignment = Alignment(horizontal="right")
                    apply_pair(ws_all, r_i, col, col + 1, a, b)
                else:
                    for jj in (col, col + 1):
                        ws_all.cell(row=r_i, column=jj, value=None).border = border
                col += 2
            r_i += 1

    ws_all.freeze_panes = "A2"
    ws_all.column_dimensions["A"].width = 8
    ws_all.column_dimensions["B"].width = 18
    for j in range(3, len(heads) + 1):
        ws_all.column_dimensions[get_column_letter(j)].width = 16

    # Mỗi thuốc 1 sheet
    for drug in ["AZI", "CIP", "TE"]:
        make_sheet(wb, drug)

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    print(f"Wrote {OUT_XLSX}")


if __name__ == "__main__":
    main()

