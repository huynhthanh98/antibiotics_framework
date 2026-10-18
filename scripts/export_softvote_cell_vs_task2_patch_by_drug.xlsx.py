#!/usr/bin/env python3
"""
Tạo 1 file Excel (xlsx) theo từng thuốc để so sánh:
- Task1 (cells)  = "Cell-level" (pipeline: Cell-level)
- Task2 (patches)= "Patch-based" (pipeline: Patch-based)

Nguồn: log trong
  - logs_all_models_cells/task1_{DRUG}_{BACKBONE}.txt
  - logs_all_models_task2/task2_{DRUG}_{BACKBONE}.txt

Trích phần: "TỔNG HỢP 5 SEEDS" -> "[Image-Level (Soft Vote)]"
Các cột theo hình: Recall(R), Precision(R), F1(R), Recall(S), Acc., AUC (mean ± std)

Output: reports/softvote_cell_vs_task2_patch_by_drug.xlsx
"""

from __future__ import annotations

import re
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
DIR_T1 = ROOT / "logs_all_models_cells"
DIR_T2 = ROOT / "logs_all_models_task2"
OUT_XLSX = ROOT / "reports" / "softvote_cell_vs_task2_patch_by_drug.xlsx"

DRUGS = ["AZI", "CIP", "TE"]

BLOCK = "Image-Level (Soft Vote)"

# Nếu muốn giới hạn backbone (như screenshot), set list này.
# Nếu None -> xuất tất cả backbone có log ở cả Task1 và Task2.
SELECTED_BACKBONES = [
    "convnext_tiny",     # ConvNeXt-Tiny
    "convnextv2_tiny",   # ConvNeXtV2-Tiny
    "densenet169",       # DenseNet-169
    "efficientnet_b3",   # EfficientNet-B3
    "resnet50",          # ResNet-50
    "vgg19",             # VGG-19
]

# key -> (label, regex stem)
METRICS = [
    ("recall_r", "Recall(R)", r"Recall\s+\(Resistant\)"),
    ("prec_r", "Precision(R)", r"Precision\s+\(Resistant\)"),
    ("f1_r", "F1(R)", r"F1\s+\(Resistant\)"),
    ("recall_s", "Recall(S)", r"Recall\s+\(Sensitive\)"),
    ("acc", "Acc.", r"Accuracy"),
    ("auc", "AUC", r"AUC"),
]


def parse_softvote_summary(path: Path) -> dict[str, tuple[float, float]] | None:
    """
    Return dict metric_key -> (mean, std) from summary block.
    """
    text = path.read_text(encoding="utf-8", errors="replace")
    if "TỔNG HỢP 5 SEEDS" not in text:
        return None

    # Extract the specific summary block for Soft Vote
    # We look for: "\n  [Image-Level (Soft Vote)]\n ... (until next block or end)"
    header = re.escape(f"[{BLOCK}]")
    m = re.search(
        rf"\n  {header}\s*\n(.*?)(?=\n  \[|\n════════════════|\Z)",
        text,
        re.DOTALL,
    )
    if not m:
        return None
    block = m.group(1)

    out: dict[str, tuple[float, float]] = {}
    for key, _, stem in METRICS:
        mm = re.search(rf"{stem}\s*:\s*([\d.]+)\s*±\s*([\d.]+)", block)
        if not mm:
            return None
        out[key] = (float(mm.group(1)), float(mm.group(2)))
    return out


def list_backbones(drug: str) -> list[str]:
    """
    backbones that exist in both task1 and task2 logs for a drug.
    """
    t1 = {p.stem.split("_", 2)[2] for p in DIR_T1.glob(f"task1_{drug}_*.txt") if not p.name.endswith("_all_models.txt")}
    t2 = {p.stem.split("_", 2)[2] for p in DIR_T2.glob(f"task2_{drug}_*.txt") if not p.name.endswith("_all_models.txt")}
    common = sorted(t1 & t2)
    if not SELECTED_BACKBONES:
        return common
    wanted = [b for b in SELECTED_BACKBONES if b in common]
    return wanted


def fmt_pm(mean: float, std: float) -> str:
    return f"{mean:.3f} ± {std:.3f}"


def main() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    fill_cell_win = PatternFill("solid", fgColor="DBEAFE")  # xanh dương
    fill_patch_win = PatternFill("solid", fgColor="DCFCE7")  # xanh lá
    fill_tie = PatternFill("solid", fgColor="FEF9C3")  # vàng
    fill_lose = PatternFill("solid", fgColor="F1F5F9")  # xám
    font_win = Font(bold=True)
    font_lose = Font(color="64748B")
    font_hdr = Font(bold=True)
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def apply_pair(ws, row_i: int, col_cell: int, col_patch: int, a: float, b: float) -> None:
        c1 = ws.cell(row=row_i, column=col_cell)
        c2 = ws.cell(row=row_i, column=col_patch)
        if abs(a - b) < 1e-12:
            for c in (c1, c2):
                c.fill = fill_tie
                c.font = Font(bold=True)
        elif a > b:
            c1.fill, c1.font = fill_cell_win, font_win
            c2.fill, c2.font = fill_lose, font_lose
        else:
            c2.fill, c2.font = fill_patch_win, font_win
            c1.fill, c1.font = fill_lose, font_lose

    wb = Workbook()
    wb.remove(wb.active)

    for drug in DRUGS:
        backbones = list_backbones(drug)
        ws = wb.create_sheet(title=drug)
        written = 0

        headers = ["Backbone", "Pipeline"]
        for _, label, _ in METRICS:
            headers.append(label)

        for j, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=j, value=h)
            cell.font = font_hdr
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

        r = 2
        for backbone in backbones:
            p1 = DIR_T1 / f"task1_{drug}_{backbone}.txt"
            p2 = DIR_T2 / f"task2_{drug}_{backbone}.txt"
            m1 = parse_softvote_summary(p1)
            m2 = parse_softvote_summary(p2)
            if not m1 or not m2:
                continue

            # Row 1: Cell-level
            ws.cell(row=r, column=1, value=backbone).border = border
            ws.cell(row=r, column=2, value="Cell-level").border = border
            ws.cell(row=r, column=1).alignment = Alignment(horizontal="left")
            ws.cell(row=r, column=2).alignment = Alignment(horizontal="left")

            # Row 2: Patch-based (Task2)
            ws.cell(row=r + 1, column=1, value=backbone).border = border
            ws.cell(row=r + 1, column=2, value="Patch-based").border = border
            ws.cell(row=r + 1, column=1).alignment = Alignment(horizontal="left")
            ws.cell(row=r + 1, column=2).alignment = Alignment(horizontal="left")

            col = 3
            for key, _, _ in METRICS:
                a_mean, a_std = m1[key]
                b_mean, b_std = m2[key]

                c_cell = ws.cell(row=r, column=col, value=fmt_pm(a_mean, a_std))
                c_patch = ws.cell(row=r + 1, column=col, value=fmt_pm(b_mean, b_std))
                for c in (c_cell, c_patch):
                    c.border = border
                    c.alignment = Alignment(horizontal="center")
                col += 1

            # Tô màu theo mean (Cells vs Patches) cho từng cột metric
            col = 3
            for key, _, _ in METRICS:
                a_mean, _ = m1[key]
                b_mean, _ = m2[key]
                c1 = ws.cell(row=r, column=col)
                c2 = ws.cell(row=r + 1, column=col)
                if abs(a_mean - b_mean) < 1e-12:
                    for c in (c1, c2):
                        c.fill = fill_tie
                        c.font = Font(bold=True)
                elif a_mean > b_mean:
                    c1.fill, c1.font = fill_cell_win, font_win
                    c2.fill, c2.font = fill_lose, font_lose
                else:
                    c2.fill, c2.font = fill_patch_win, font_win
                    c1.fill, c1.font = fill_lose, font_lose
                col += 1

            # visual group border: make backbone cell merged? keep repeated like screenshot (2 rows)
            r += 2
            written += 1

        # column widths
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 12
        for j in range(3, len(headers) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 16
        ws.freeze_panes = "A2"
        print(f"[{drug}] backbones matched={len(backbones)} | rows written={written}")

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    print(f"Wrote {OUT_XLSX}")


if __name__ == "__main__":
    main()

