#!/usr/bin/env python3
"""
Sinh Figure 1 (bar plot Recall), Figure 2 (heatmap p-value),
Figure 3 (forest plot ΔRecall) và Table 1 (summary CSV/xlsx)
so sánh Cell-level vs Patch-based trên 12 backbones × 3 kháng sinh.

Nguồn: reports/softvote_cell_vs_patch_rpf1_auc.xlsx
Output:
  reports/table1_summary_cell_vs_patch.csv
  reports/table1_summary_cell_vs_patch.xlsx
  reports/fig1_recall_bar_cell_vs_patch.pdf / .png
  reports/fig2_pval_heatmap_cell_vs_patch.pdf / .png
  reports/fig3_forest_delta_recall.pdf / .png
"""

from __future__ import annotations

from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import matplotlib.ticker as ticker
import numpy as np
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side, numbers
)
from openpyxl.utils import get_column_letter

# ── Paths ────────────────────────────────────────────────────────────────────
ROOT = Path(__file__).resolve().parents[1]
XLSX_IN  = ROOT / "reports" / "softvote_cell_vs_patch_rpf1_auc.xlsx"
REPORT   = ROOT / "reports"

DRUGS = ["AZI", "CIP", "TE"]
ALPHA = 0.05

# ── Backbone display names ───────────────────────────────────────────────────
BACKBONE_LABELS = {
    "resnet50":       "ResNet-50",
    "resnet101":      "ResNet-101",
    "vgg16":          "VGG-16",
    "vgg19":          "VGG-19",
    "swin_tiny":      "Swin-T",
    "swin_small":     "Swin-S",
    "efficientnet_b3":"EffNet-B3",
    "efficientnet_b5":"EffNet-B5",
    "densenet121":    "DenseNet-121",
    "densenet201":    "DenseNet-201",
    "convnext_tiny":  "ConvNeXt-T",
    "convnextv2_tiny":"ConvNeXtV2-T",
}

# ── Colors ───────────────────────────────────────────────────────────────────
COLOR_CELL  = "#2563EB"   # blue-600
COLOR_PATCH = "#F97316"   # orange-500
COLOR_SIG   = "#16A34A"   # green-600
COLOR_NS    = "#DC2626"   # red-600

DRUG_COLORS = {"AZI": "#1D4ED8", "CIP": "#7C3AED", "TE": "#B45309"}


# ═════════════════════════════════════════════════════════════════════════════
#  DATA LOADING
# ═════════════════════════════════════════════════════════════════════════════

def load_data(xlsx_path: Path) -> dict:
    """
    Returns dict keyed by drug → list of dicts per backbone:
      {backbone, cell_recall, cell_recall_std, patch_recall, patch_recall_std,
       cell_auc, cell_auc_std, patch_auc, patch_auc_std, p_recall, p_auc}
    """
    wb = openpyxl.load_workbook(xlsx_path)
    data = {}

    for drug in DRUGS:
        ws = wb[drug]
        rows = list(ws.iter_rows(values_only=True))[1:]  # skip header

        cell_map  = {}
        patch_map = {}
        for row in rows:
            bb, pipeline = row[0], row[1]
            recall_str = row[2]  # "0.910 ± 0.035"
            auc_str    = row[5]
            p_recall   = row[6]
            p_auc      = row[7]

            mean_r, std_r = _parse_pm(recall_str)
            mean_a, std_a = _parse_pm(auc_str)

            if pipeline == "Cell-level":
                cell_map[bb] = dict(
                    recall=mean_r, recall_std=std_r,
                    auc=mean_a,    auc_std=std_a,
                    p_recall=p_recall, p_auc=p_auc,
                )
            elif "Patch" in pipeline:
                patch_map[bb] = dict(
                    recall=mean_r, recall_std=std_r,
                    auc=mean_a,    auc_std=std_a,
                )

        records = []
        for bb in cell_map:
            if bb not in patch_map:
                continue
            c = cell_map[bb]
            p = patch_map[bb]
            records.append({
                "backbone":         bb,
                "cell_recall":      c["recall"],
                "cell_recall_std":  c["recall_std"],
                "patch_recall":     p["recall"],
                "patch_recall_std": p["recall_std"],
                "cell_auc":         c["auc"],
                "cell_auc_std":     c["auc_std"],
                "patch_auc":        p["auc"],
                "patch_auc_std":    p["auc_std"],
                "p_recall":         c["p_recall"],
                "p_auc":            c["p_auc"],
                "delta_recall":     c["recall"] - p["recall"],
            })
        data[drug] = records

    return data


def _parse_pm(s: str) -> tuple[float, float]:
    parts = s.split("±")
    return float(parts[0].strip()), float(parts[1].strip())


# ═════════════════════════════════════════════════════════════════════════════
#  TABLE 1  — summary CSV + XLSX
# ═════════════════════════════════════════════════════════════════════════════

def make_table1(data: dict) -> None:
    rows_csv = []
    for drug in DRUGS:
        recs = data[drug]
        # aggregate across backbones
        cell_recalls  = [r["cell_recall"]  for r in recs]
        patch_recalls = [r["patch_recall"] for r in recs]
        cell_aucs     = [r["cell_auc"]     for r in recs]
        patch_aucs    = [r["patch_auc"]    for r in recs]
        p_recalls     = [r["p_recall"]     for r in recs]
        n_sig         = sum(1 for p in p_recalls if p is not None and p < ALPHA)
        n_total       = len(recs)
        delta_mean    = np.mean([r["delta_recall"] for r in recs])

        rows_csv.append({
            "Drug": drug, "Pipeline": "Cell-level",
            "Recall(R) mean": f"{np.mean(cell_recalls):.3f}",
            "Recall(R) std":  f"{np.std(cell_recalls):.3f}",
            "AUC mean":       f"{np.mean(cell_aucs):.3f}",
            "AUC std":        f"{np.std(cell_aucs):.3f}",
            "Δ Recall (Cell−Patch)": f"+{delta_mean:.3f}",
            "# sig. backbone (p<0.05)": f"{n_sig}/{n_total}",
        })
        rows_csv.append({
            "Drug": drug, "Pipeline": "Patch-based",
            "Recall(R) mean": f"{np.mean(patch_recalls):.3f}",
            "Recall(R) std":  f"{np.std(patch_recalls):.3f}",
            "AUC mean":       f"{np.mean(patch_aucs):.3f}",
            "AUC std":        f"{np.std(patch_aucs):.3f}",
            "Δ Recall (Cell−Patch)": "—",
            "# sig. backbone (p<0.05)": "—",
        })

    # CSV
    import csv
    csv_path = REPORT / "table1_summary_cell_vs_patch.csv"
    fieldnames = list(rows_csv[0].keys())
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows_csv)
    print(f"Wrote {csv_path}")

    # XLSX
    wb2 = Workbook()
    ws = wb2.active
    ws.title = "Table1"
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    font_hdr = Font(bold=True, color="FFFFFF")
    fill_hdr = PatternFill("solid", fgColor="1E40AF")  # blue-800
    fill_cell_row = PatternFill("solid", fgColor="DBEAFE")   # blue-100
    fill_patch_row = PatternFill("solid", fgColor="FEF3C7")  # amber-100
    fill_drug_sep = PatternFill("solid", fgColor="F1F5F9")   # slate-100

    headers = fieldnames
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = font_hdr
        c.fill = fill_hdr
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, row_data in enumerate(rows_csv, 2):
        is_cell = row_data["Pipeline"] == "Cell-level"
        fill = fill_cell_row if is_cell else fill_patch_row
        for j, key in enumerate(headers, 1):
            c = ws.cell(row=i, column=j, value=row_data[key])
            c.border = border
            c.fill = fill
            c.alignment = Alignment(horizontal="center")
            if j == 1 or j == 2:
                c.alignment = Alignment(horizontal="left")
            if is_cell and key == "# sig. backbone (p<0.05)":
                sig_str = row_data[key]
                if "/" in sig_str:
                    num, den = sig_str.split("/")
                    if int(num) >= 8:
                        c.font = Font(bold=True, color="166534")  # green
            if is_cell and key == "Δ Recall (Cell−Patch)":
                c.font = Font(bold=True, color="1D4ED8")

    col_widths = [8, 14, 16, 14, 12, 10, 22, 24]
    for j, w_ in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w_
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "A2"

    xlsx_path = REPORT / "table1_summary_cell_vs_patch.xlsx"
    wb2.save(xlsx_path)
    print(f"Wrote {xlsx_path}")


# ═════════════════════════════════════════════════════════════════════════════
#  FIGURE 1  — Grouped bar plot: Recall(R) Cell vs Patch (3 subplots)
# ═════════════════════════════════════════════════════════════════════════════

def make_fig1_bar(data: dict) -> None:
    fig, axes = plt.subplots(3, 1, figsize=(15, 13), constrained_layout=True)
    fig.suptitle(
        "Recall (Resistant class): Cell-level vs Patch-based\n"
        "Mean ± std over 5 seeds, 12 backbone architectures",
        fontsize=14, fontweight="bold", y=1.01,
    )

    for ax, drug in zip(axes, DRUGS):
        recs = data[drug]
        n = len(recs)
        x = np.arange(n)
        w = 0.35

        cell_means  = np.array([r["cell_recall"]      for r in recs])
        cell_stds   = np.array([r["cell_recall_std"]  for r in recs])
        patch_means = np.array([r["patch_recall"]      for r in recs])
        patch_stds  = np.array([r["patch_recall_std"] for r in recs])
        p_vals      = [r["p_recall"] for r in recs]
        bb_labels   = [BACKBONE_LABELS.get(r["backbone"], r["backbone"]) for r in recs]

        bars_cell = ax.bar(
            x - w/2, cell_means, w,
            yerr=cell_stds, capsize=4,
            color=COLOR_CELL, alpha=0.85,
            label="Cell-level",
            error_kw=dict(elinewidth=1.2, ecolor="#1E3A8A"),
        )
        bars_patch = ax.bar(
            x + w/2, patch_means, w,
            yerr=patch_stds, capsize=4,
            color=COLOR_PATCH, alpha=0.85,
            label="Patch-based",
            error_kw=dict(elinewidth=1.2, ecolor="#7C2D12"),
        )

        # Significance annotation
        y_max = np.maximum(cell_means + cell_stds, patch_means + patch_stds)
        for i, (p, ym) in enumerate(zip(p_vals, y_max)):
            sig_y = min(ym + 0.03, 1.02)
            if p is not None and p < ALPHA:
                ax.text(x[i], sig_y, "*", ha="center", va="bottom",
                        fontsize=14, color=COLOR_SIG, fontweight="bold")
            else:
                ax.text(x[i], sig_y, "n.s.", ha="center", va="bottom",
                        fontsize=7, color="#94A3B8")

        # Delta recall annotation on top of cell bar
        for i, (cm, pm) in enumerate(zip(cell_means, patch_means)):
            delta = cm - pm
            ax.text(x[i] - w/2, cm + cell_stds[i] + 0.045,
                    f"+{delta:.2f}", ha="center", va="bottom",
                    fontsize=6.5, color=COLOR_CELL, style="italic")

        ax.set_xlim(-0.6, n - 0.4)
        ax.set_ylim(0, 1.13)
        ax.set_xticks(x)
        ax.set_xticklabels(bb_labels, rotation=30, ha="right", fontsize=9)
        ax.set_ylabel("Recall (Resistant class)", fontsize=10)
        ax.yaxis.set_major_formatter(ticker.FormatStrFormatter("%.2f"))
        ax.axhline(1.0, color="gray", linestyle="--", linewidth=0.7, alpha=0.5)
        ax.grid(axis="y", linestyle="--", linewidth=0.5, alpha=0.4)
        ax.spines[["top", "right"]].set_visible(False)

        # Drug label badge
        n_sig = sum(1 for p in p_vals if p is not None and p < ALPHA)
        cell_avg  = np.mean(cell_means)
        patch_avg = np.mean(patch_means)
        delta_avg = cell_avg - patch_avg
        ax.set_title(
            f"{drug}  |  Cell: {cell_avg:.3f}   Patch: {patch_avg:.3f}   "
            f"Δ={delta_avg:+.3f}   Sig: {n_sig}/12",
            fontsize=11, fontweight="bold", loc="left", pad=6,
            color=DRUG_COLORS[drug],
        )
        ax.legend(loc="lower right", fontsize=9, framealpha=0.85)

    # Shared legend note
    star_patch = mpatches.Patch(color=COLOR_SIG,  label="* p < 0.05 (one-sided t-test)")
    ns_patch   = mpatches.Patch(color="#94A3B8",  label="n.s. p ≥ 0.05")
    fig.legend(
        handles=[star_patch, ns_patch],
        loc="lower center", ncol=2, fontsize=9,
        bbox_to_anchor=(0.5, -0.02), framealpha=0.9,
    )

    for fmt in ("pdf", "png"):
        out = REPORT / f"fig1_recall_bar_cell_vs_patch.{fmt}"
        fig.savefig(out, dpi=180, bbox_inches="tight")
        print(f"Wrote {out}")
    plt.close(fig)


# ═════════════════════════════════════════════════════════════════════════════
#  FIGURE 2  — Heatmap p-value Recall (backbone × drug)
# ═════════════════════════════════════════════════════════════════════════════

def make_fig2_heatmap(data: dict) -> None:
    # Collect backbone order (use AZI order as canonical)
    backbone_order = [r["backbone"] for r in data["AZI"]]
    n_bb = len(backbone_order)
    n_drug = len(DRUGS)

    pval_matrix = np.full((n_bb, n_drug), np.nan)
    delta_matrix = np.full((n_bb, n_drug), np.nan)

    for j, drug in enumerate(DRUGS):
        rec_map = {r["backbone"]: r for r in data[drug]}
        for i, bb in enumerate(backbone_order):
            if bb in rec_map:
                pval_matrix[i, j]  = rec_map[bb]["p_recall"]
                delta_matrix[i, j] = rec_map[bb]["delta_recall"]

    fig, axes = plt.subplots(1, 2, figsize=(13, 7), constrained_layout=True)
    fig.suptitle(
        "Statistical comparison: Cell-level vs Patch-based\n"
        "Paired one-sided t-test  (H₁: Cell > Patch, α = 0.05)",
        fontsize=13, fontweight="bold",
    )

    # ── Left: p-value heatmap ─────────────────────────────────────────────
    ax = axes[0]
    from matplotlib.colors import LinearSegmentedColormap
    cmap_pval = LinearSegmentedColormap.from_list(
        "pval_cmap",
        [(0.0, "#166534"), (0.05, "#BBF7D0"), (0.06, "#FEE2E2"), (1.0, "#991B1B")],
    )
    im = ax.imshow(pval_matrix, cmap=cmap_pval, vmin=0, vmax=1, aspect="auto")
    plt.colorbar(im, ax=ax, label="p-value", fraction=0.046, pad=0.04)

    ax.set_xticks(range(n_drug))
    ax.set_xticklabels(DRUGS, fontsize=11, fontweight="bold")
    ax.set_yticks(range(n_bb))
    ax.set_yticklabels([BACKBONE_LABELS.get(b, b) for b in backbone_order], fontsize=9)
    ax.set_title("p-value (Recall, one-sided)", fontsize=11, pad=10)

    for i in range(n_bb):
        for j in range(n_drug):
            p = pval_matrix[i, j]
            if np.isnan(p):
                continue
            sig_marker = "★" if p < ALPHA else ""
            txt = f"{p:.3f}{sig_marker}"
            txt_color = "white" if p < ALPHA else "#1F2937"
            fontw = "bold" if p < ALPHA else "normal"
            ax.text(j, i, txt, ha="center", va="center",
                    fontsize=8, color=txt_color, fontweight=fontw)

    # Threshold line annotation
    ax.axhline(-0.5, color="white", linewidth=0)
    ax.set_xlabel("Antibiotic", fontsize=10)

    # ── Right: ΔRecall heatmap ────────────────────────────────────────────
    ax2 = axes[1]
    cmap_delta = LinearSegmentedColormap.from_list(
        "delta_cmap",
        [(0.0, "#7C2D12"), (0.3, "#FED7AA"), (0.5, "#F0FDF4"), (1.0, "#14532D")],
    )
    im2 = ax2.imshow(delta_matrix, cmap=cmap_delta, vmin=-0.1, vmax=0.5, aspect="auto")
    plt.colorbar(im2, ax=ax2, label="Δ Recall (Cell − Patch)", fraction=0.046, pad=0.04)

    ax2.set_xticks(range(n_drug))
    ax2.set_xticklabels(DRUGS, fontsize=11, fontweight="bold")
    ax2.set_yticks(range(n_bb))
    ax2.set_yticklabels([BACKBONE_LABELS.get(b, b) for b in backbone_order], fontsize=9)
    ax2.set_title("Δ Recall = Cell − Patch", fontsize=11, pad=10)

    for i in range(n_bb):
        for j in range(n_drug):
            delta = delta_matrix[i, j]
            p     = pval_matrix[i, j]
            if np.isnan(delta):
                continue
            sig_marker = "★" if (p is not None and not np.isnan(p) and p < ALPHA) else ""
            txt = f"{delta:+.3f}{sig_marker}"
            txt_color = "white" if delta > 0.25 else "#1F2937"
            fontw = "bold" if (not np.isnan(p) and p < ALPHA) else "normal"
            ax2.text(j, i, txt, ha="center", va="center",
                     fontsize=8, color=txt_color, fontweight=fontw)

    ax2.set_xlabel("Antibiotic", fontsize=10)

    # Legend for ★
    fig.text(
        0.5, -0.01,
        "★ = p < 0.05 (one-sided t-test, H₁: Cell > Patch)",
        ha="center", fontsize=10, style="italic", color="#374151",
    )

    for fmt in ("pdf", "png"):
        out = REPORT / f"fig2_pval_heatmap_cell_vs_patch.{fmt}"
        fig.savefig(out, dpi=180, bbox_inches="tight")
        print(f"Wrote {out}")
    plt.close(fig)


# ═════════════════════════════════════════════════════════════════════════════
#  FIGURE 3  — Forest / dot plot: ΔRecall per backbone × drug
# ═════════════════════════════════════════════════════════════════════════════

def make_fig3_forest(data: dict) -> None:
    backbone_order = [r["backbone"] for r in data["AZI"]]
    n_bb = len(backbone_order)
    bb_labels = [BACKBONE_LABELS.get(b, b) for b in backbone_order]

    fig, axes = plt.subplots(1, 3, figsize=(15, 6), sharey=True, constrained_layout=True)
    fig.suptitle(
        "Δ Recall (Cell − Patch) per backbone architecture\n"
        "★ = p < 0.05 (one-sided t-test, H₁: Cell > Patch)",
        fontsize=13, fontweight="bold",
    )

    for ax, drug in zip(axes, DRUGS):
        rec_map = {r["backbone"]: r for r in data[drug]}
        deltas = []
        colors = []
        markers = []
        for bb in backbone_order:
            r = rec_map.get(bb)
            if r:
                deltas.append(r["delta_recall"])
                p = r["p_recall"]
                sig = p is not None and p < ALPHA
                colors.append(COLOR_SIG if sig else "#94A3B8")
                markers.append("★" if sig else "○")
            else:
                deltas.append(np.nan)
                colors.append("gray")
                markers.append("")

        y = np.arange(n_bb)
        ax.axvline(0, color="#374151", linewidth=1.2, linestyle="--", alpha=0.6)

        for i, (d, c) in enumerate(zip(deltas, colors)):
            if not np.isnan(d):
                ax.barh(y[i], d, height=0.55, color=c, alpha=0.80, zorder=3)
                # label value
                ha = "left" if d >= 0 else "right"
                offset = 0.005 if d >= 0 else -0.005
                ax.text(d + offset, y[i], f"{d:+.3f}", va="center", ha=ha,
                        fontsize=7.5, color=c, fontweight="bold" if c == COLOR_SIG else "normal")

        ax.set_yticks(y)
        ax.set_yticklabels(bb_labels, fontsize=9)
        ax.set_xlabel("Δ Recall (Cell − Patch)", fontsize=10)
        ax.set_xlim(-0.15, 0.62)
        ax.xaxis.set_major_formatter(ticker.FormatStrFormatter("%+.2f"))
        ax.grid(axis="x", linestyle="--", linewidth=0.5, alpha=0.4)
        ax.spines[["top", "right"]].set_visible(False)

        n_sig = sum(1 for d, c in zip(deltas, colors)
                    if not np.isnan(d) and c == COLOR_SIG)
        mean_d = np.nanmean(deltas)
        ax.set_title(
            f"{drug}  |  mean Δ={mean_d:+.3f}\n{n_sig}/12 sig. (p<0.05)",
            fontsize=11, fontweight="bold", color=DRUG_COLORS[drug],
        )
        ax.invert_yaxis()

    sig_patch = mpatches.Patch(color=COLOR_SIG, label="p < 0.05 (significant)")
    ns_patch  = mpatches.Patch(color="#94A3B8", label="p ≥ 0.05 (not significant)")
    fig.legend(
        handles=[sig_patch, ns_patch],
        loc="lower center", ncol=2, fontsize=9,
        bbox_to_anchor=(0.5, -0.03), framealpha=0.9,
    )

    for fmt in ("pdf", "png"):
        out = REPORT / f"fig3_forest_delta_recall.{fmt}"
        fig.savefig(out, dpi=180, bbox_inches="tight")
        print(f"Wrote {out}")
    plt.close(fig)


# ═════════════════════════════════════════════════════════════════════════════
#  MAIN
# ═════════════════════════════════════════════════════════════════════════════

def main() -> None:
    print(f"Loading {XLSX_IN}")
    data = load_data(XLSX_IN)

    print("\n── Table 1 ──")
    make_table1(data)

    print("\n── Figure 1: Bar plot ──")
    make_fig1_bar(data)

    print("\n── Figure 2: Heatmap ──")
    make_fig2_heatmap(data)

    print("\n── Figure 3: Forest plot ──")
    make_fig3_forest(data)

    print("\nDone. All outputs in:", REPORT)


if __name__ == "__main__":
    main()
