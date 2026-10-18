#!/usr/bin/env python3
"""
So sánh metric (nhãn Resistant) giữa log Task1 (cells) và Task2 (patches)
từ thư mục logs_all_models_cells và logs_all_models_task2.
Đọc khối 'TỔNG HỢP 5 SEEDS' (mean ± std); so sánh theo mean.
"""
from __future__ import annotations

import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
CELLS_DIR = ROOT / "logs_all_models_cells"
TASK2_DIR = ROOT / "logs_all_models_task2"
OUT_HTML = ROOT / "reports" / "task1_cells_vs_task2_patches_resistant.html"
OUT_CSV = ROOT / "reports" / "task1_vs_task2_resistant_metrics.csv"
OUT_XLSX = ROOT / "reports" / "task1_vs_task2_resistant_metrics.xlsx"

LEVELS = [
    ("Cell-Level", "Cell-Level"),
    ("Image-Level (Soft Vote)", "Image-Level (Soft Vote)"),
    ("Image-Level (Majority Vote)", "Image-Level (Majority Vote)"),
]

METRICS = [
    ("recall_r", "Recall(R)"),
    ("prec_r", "Precision(R)"),
    ("f1_r", "F1(R)"),
    ("acc", "Accuracy"),
    ("auc", "AUC"),
]


def parse_summary_block(text: str, block_title: str) -> dict[str, float] | None:
    block_re = (
        rf"\n  \[{re.escape(block_title)}\]\s*\n"
        r"(.*?)(?=\n  \[|\n════════════════|\Z)"
    )
    bm = re.search(block_re, text, re.DOTALL)
    if not bm:
        return None
    block = bm.group(1)
    metrics: dict[str, float] = {}
    for line in block.splitlines():
        if re.match(r"\s+Recall\s+\(Resistant\)\s*:\s*([\d.]+)", line):
            metrics["recall_r"] = float(
                re.match(r"\s+Recall\s+\(Resistant\)\s*:\s*([\d.]+)", line).group(1)
            )
        elif re.match(r"\s+Precision\s+\(Resistant\)\s*:\s*([\d.]+)", line):
            metrics["prec_r"] = float(
                re.match(r"\s+Precision\s+\(Resistant\)\s*:\s*([\d.]+)", line).group(1)
            )
        elif re.match(r"\s+F1\s+\(Resistant\)\s*:\s*([\d.]+)", line):
            metrics["f1_r"] = float(
                re.match(r"\s+F1\s+\(Resistant\)\s*:\s*([\d.]+)", line).group(1)
            )
        elif re.match(r"\s+Accuracy\s+\s*:\s*([\d.]+)", line):
            metrics["acc"] = float(
                re.match(r"\s+Accuracy\s+\s*:\s*([\d.]+)", line).group(1)
            )
        elif re.match(r"\s+AUC\s+\s*:\s*([\d.]+)", line):
            metrics["auc"] = float(re.match(r"\s+AUC\s+\s*:\s*([\d.]+)", line).group(1))
    if len(metrics) != 5:
        return None
    return metrics


def parse_log(path: Path) -> dict | None:
    text = path.read_text(encoding="utf-8", errors="replace")
    if "TỔNG HỢP 5 SEEDS" not in text:
        return None
    out: dict = {}
    for block_title, key in LEVELS:
        m = parse_summary_block(text, block_title)
        out[key] = m
    return out


def collect(prefix: str, folder: Path) -> dict[tuple[str, str], dict]:
    out = {}
    for p in sorted(folder.glob(f"{prefix}_*.txt")):
        if p.name.endswith("_all_models.txt") or p.name == "status.txt":
            continue
        rest = p.stem[len(prefix) + 1 :]
        if "_" not in rest:
            continue
        drug, backbone = rest.split("_", 1)
        parsed = parse_log(p)
        if not parsed or not any(parsed.values()):
            continue
        out[(drug, backbone)] = parsed
    return out


def style_t1(a: float, b: float) -> str:
    """Ô giá trị Task1: nổi bật nếu Task1 cao hơn."""
    if abs(a - b) < 1e-9:
        return "background:#fef9c3;font-weight:600;"
    if a > b:
        return "background:#dbeafe;font-weight:700;"
    return "background:#f1f5f9;color:#64748b;"


def style_t2(a: float, b: float) -> str:
    """Ô giá trị Task2: nổi bật nếu Task2 cao hơn."""
    if abs(a - b) < 1e-9:
        return "background:#fef9c3;font-weight:600;"
    if b > a:
        return "background:#dcfce7;font-weight:700;"
    return "background:#f1f5f9;color:#64748b;"


def fmt(x: float) -> str:
    return f"{x:.4f}"


def write_xlsx(
    t1: dict,
    t2: dict,
    keys: list,
    rows_flat: list[tuple[str, str, str, dict, dict]],
) -> None:
    """rows_flat: (drug, level_key, backbone, m1, m2) for All_data sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    fill_t1_win = PatternFill("solid", fgColor="DBEAFE")
    fill_t2_win = PatternFill("solid", fgColor="DCFCE7")
    fill_tie = PatternFill("solid", fgColor="FEF9C3")
    fill_lose = PatternFill("solid", fgColor="F1F5F9")
    font_win = Font(bold=True)
    font_lose = Font(color="64748B")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True)

    def apply_pair(ws, row: int, c_t1: int, c_t2: int, a: float, b: float) -> None:
        if abs(a - b) < 1e-9:
            for c in (c_t1, c_t2):
                cell = ws.cell(row=row, column=c)
                cell.fill = fill_tie
                cell.font = Font(bold=True)
        else:
            c1 = ws.cell(row=row, column=c_t1)
            c2 = ws.cell(row=row, column=c_t2)
            if a > b:
                c1.fill, c1.font = fill_t1_win, font_win
                c2.fill, c2.font = fill_lose, font_lose
            else:
                c2.fill, c2.font = fill_t2_win, font_win
                c1.fill, c1.font = fill_lose, font_lose

    wb = Workbook()
    # Một sheet phẳng: toàn bộ drug × level × backbone
    ws_all = wb.active
    ws_all.title = "All_data"
    hdr = ["drug", "level", "backbone"] + [
        f"t1_{k}" for k, _ in METRICS
    ] + [f"t2_{k}" for k, _ in METRICS]
    for col, h in enumerate(hdr, start=1):
        cell = ws_all.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.border = border
    for r, (drug, level_key, backbone, m1, m2) in enumerate(rows_flat, start=2):
        ws_all.cell(row=r, column=1, value=drug).border = border
        ws_all.cell(row=r, column=2, value=level_key).border = border
        ws_all.cell(row=r, column=3, value=backbone).border = border
        col = 4
        for key, _ in METRICS:
            ws_all.cell(row=r, column=col, value=m1[key]).border = border
            col += 1
        for key, _ in METRICS:
            ws_all.cell(row=r, column=col, value=m2[key]).border = border
            col += 1
        # Cột 4–8 = t1_*, 9–13 = t2_* → cặp (4,9), (5,10), ...
        for i, (key, _) in enumerate(METRICS):
            apply_pair(ws_all, r, 4 + i, 9 + i, m1[key], m2[key])
        for c in range(1, len(hdr) + 1):
            ws_all.cell(row=r, column=c).alignment = Alignment(horizontal="right")
        ws_all.cell(row=r, column=1).alignment = Alignment(horizontal="left")
        ws_all.cell(row=r, column=2).alignment = Alignment(horizontal="left")
        ws_all.cell(row=r, column=3).alignment = Alignment(horizontal="left")

    for c in range(1, len(hdr) + 1):
        ws_all.column_dimensions[get_column_letter(c)].width = 14
    ws_all.column_dimensions["C"].width = 18
    ws_all.freeze_panes = "A2"

    # Mỗi level một sheet (chỉ backbone + metric), dễ đọc
    for level_key, level_title in LEVELS:
        short = (
            level_key.replace(" ", "_")
            .replace("(", "")
            .replace(")", "")
            .replace("/", "_")[:31]
        )
        ws = wb.create_sheet(title=short)
        heads = ["Drug", "Backbone"] + [
            f"{lbl} (T1)" for _, lbl in METRICS
        ] + [f"{lbl} (T2)" for _, lbl in METRICS]
        for col, h in enumerate(heads, start=1):
            c = ws.cell(row=1, column=col, value=h)
            c.font = header_font
            c.border = border
        row = 2
        for drug in ["AZI", "CIP", "TE"]:
            for _, backbone in sorted(k for k in keys if k[0] == drug):
                m1 = t1[(drug, backbone)].get(level_key)
                m2 = t2[(drug, backbone)].get(level_key)
                if not m1 or not m2:
                    continue
                ws.cell(row=row, column=1, value=drug).border = border
                ws.cell(row=row, column=2, value=backbone).border = border
                col = 3
                for key, _ in METRICS:
                    ws.cell(row=row, column=col, value=round(m1[key], 6)).border = border
                    col += 1
                for key, _ in METRICS:
                    ws.cell(row=row, column=col, value=round(m2[key], 6)).border = border
                    col += 1
                base = 3
                for key, _ in METRICS:
                    apply_pair(ws, row, base, base + 5, m1[key], m2[key])
                    base += 1
                for c in range(1, len(heads) + 1):
                    ws.cell(row=row, column=c).alignment = Alignment(horizontal="right")
                ws.cell(row=row, column=1).alignment = Alignment(horizontal="left")
                ws.cell(row=row, column=2).alignment = Alignment(horizontal="left")
                row += 1
        for c in range(1, len(heads) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 13
        ws.column_dimensions["B"].width = 20
        ws.freeze_panes = "A2"

    wb.save(OUT_XLSX)


def main() -> None:
    t1 = collect("task1", CELLS_DIR)
    t2 = collect("task2", TASK2_DIR)
    keys = sorted(set(t1.keys()) & set(t2.keys()))

    OUT_HTML.parent.mkdir(parents=True, exist_ok=True)

    csv_lines = [
        "drug,level,backbone,"
        + ",".join([f"t1_{k},t2_{k}" for k, _ in METRICS])
    ]
    rows_for_xlsx: list[tuple[str, str, str, dict, dict]] = []

    parts = [
        "<!DOCTYPE html><html><head><meta charset='utf-8'>",
        "<title>Task1 (cells) vs Task2 (patches) — Resistant class</title>",
        "<style>",
        "body{font-family:system-ui,sans-serif;margin:24px;background:#fafafa;color:#111;}",
        "h1{font-size:1.25rem;}",
        "h2{font-size:1.05rem;margin-top:2rem;color:#334155;}",
        "h3{font-size:0.95rem;margin-top:1.25rem;color:#475569;}",
        "table{border-collapse:collapse;width:100%;max-width:1400px;background:#fff;box-shadow:0 1px 3px #0001;margin-bottom:0.5rem;font-size:12px;}",
        "th,td{border:1px solid #e2e8f0;padding:6px 8px;text-align:right;}",
        "th:first-child,td:first-child{text-align:left;min-width:140px;}",
        ".legend span{display:inline-block;padding:4px 10px;margin-right:8px;border-radius:4px;font-size:12px;}",
        ".l1{background:#dbeafe;} .l2{background:#dcfce7;} .tie{background:#fef9c3;}",
        "</style></head><body>",
        "<h1>So sánh Task1 (cells / masked) vs Task2 (patches / full slide)</h1>",
        "<p>Metric trích từ dòng <code>TỔNG HỢP 5 SEEDS</code> trong log (mean; phần sau ± bỏ qua khi so sánh).</p>",
        "<div class='legend'><span class='l1'>Task1 (cells) thắng metric</span>",
        "<span class='l2'>Task2 (patches) thắng metric</span>",
        "<span class='tie'>Bằng nhau</span> <small>(ô thua: xám nhạt)</small></div>",
    ]

    for drug in ["AZI", "CIP", "TE"]:
        dk = [k for k in keys if k[0] == drug]
        if not dk:
            continue
        parts.append(f"<h2>Drug: {drug}</h2>")
        for level_key, level_title in LEVELS:
            parts.append(f"<h3>{level_title}</h3>")
            parts.append("<table><thead><tr>")
            parts.append("<th>Backbone</th>")
            for _, label in METRICS:
                parts.append(f"<th colspan='2'>{label}<br><small>Task1 | Task2</small></th>")
            parts.append("</tr></thead><tbody>")
            for _, backbone in dk:
                m1 = t1[(drug, backbone)].get(level_key)
                m2 = t2[(drug, backbone)].get(level_key)
                if not m1 or not m2:
                    continue
                row_csv = [drug, level_key, backbone]
                parts.append("<tr>")
                parts.append(f"<td><b>{backbone}</b></td>")
                for key, _ in METRICS:
                    a, b = m1[key], m2[key]
                    row_csv.extend([f"{a:.6f}", f"{b:.6f}"])
                    parts.append(
                        f"<td style='{style_t1(a,b)}'>{fmt(a)}</td>"
                        f"<td style='{style_t2(a,b)}'>{fmt(b)}</td>"
                    )
                parts.append("</tr>")
                csv_lines.append(",".join(row_csv))
                rows_for_xlsx.append((drug, level_key, backbone, m1, m2))
            parts.append("</tbody></table>")

    parts.append("</body></html>")
    OUT_HTML.write_text("".join(parts), encoding="utf-8")

    OUT_CSV.write_text("\n".join(csv_lines), encoding="utf-8")
    write_xlsx(t1, t2, keys, rows_for_xlsx)
    print(f"Wrote {OUT_HTML}")
    print(f"Wrote {OUT_CSV}")
    print(f"Wrote {OUT_XLSX}")
    print(f"Combinations (drug, backbone): {len(keys)}")


if __name__ == "__main__":
    main()
