"""
Degraded Mask Robustness Evaluation — Task 1 models.

Mục tiêu: Kiểm tra xem model Task 1 (train trên ảnh cell đã segment sẵn)
có bền vững như thế nào khi segmentation mask không hoàn hảo.

Không train lại model — chỉ load checkpoint đã train, degrade ảnh test,
rồi đánh giá và so sánh với baseline (không degrade).

Degradation modes:
  occlusion  : Che ngẫu nhiên X% diện tích bằng patch đen
               → mô phỏng phần tế bào bị mất (under-segmentation)
  erosion    : Erode vùng tế bào vào trong N pixel từ boundary
               → mô phỏng mask boundary bị thu hẹp
  bg_noise   : Thêm Gaussian noise vào vùng background (pixel ≈ 0)
               → mô phỏng background lọt vào mask (over-segmentation)

Tái sử dụng:
  - Task1Dataset (datasets.py)
  - evaluate_loader (trainer.py)
  - compute_metrics, image_level_voting (metrics.py)
  - get_model (models.py)
  - Config (config.py)

Usage:
    # Đánh giá 1 model, 1 drug, tất cả degradation:
    python degraded_mask_eval.py --drug CIP --model efficientnet_b5

    # Nhiều models:
    python degraded_mask_eval.py --drug CIP --model efficientnet_b5 vit_small convnext_base

    # Chỉ 1 degradation mode:
    python degraded_mask_eval.py --drug CIP --model efficientnet_b5 --mode occlusion

    # Output file:
    python degraded_mask_eval.py --drug CIP --model efficientnet_b5 --out result.xlsx

    # Seeds cụ thể:
    python degraded_mask_eval.py --drug CIP --model efficientnet_b5 --seeds 0 1 2
"""

import os
import sys
import argparse
import warnings
import random
import hashlib
import numpy as np
import pandas as pd
import torch
import torch.nn as nn
from torch.utils.data import Dataset, DataLoader
from PIL import Image, ImageFilter
from pathlib import Path
from typing import Dict, List, Optional, Tuple

warnings.filterwarnings('ignore')

# ── Import framework modules ──────────────────────────────────────────────────
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from config import Config
from datasets import Task1Dataset, get_test_transform, label_to_str, _normalize_df
from models import get_model
from metrics import compute_metrics, image_level_voting, evaluate_full
from trainer import evaluate_loader


# =============================================================================
# Degradation functions
# =============================================================================

def _apply_random_occlusion(
    img: Image.Image, level: float, rng: np.random.RandomState
) -> Image.Image:
    """
    Che ngẫu nhiên 'level' tỉ lệ diện tích bằng các hình chữ nhật đen.
    Mô phỏng: phần tế bào bị mất do segmentation under-detect.

    level: 0.0 (không che) → 0.5 (che 50% diện tích)
    """
    if level <= 0:
        return img
    arr = np.array(img)
    H, W = arr.shape[:2]
    target_pixels = int(H * W * level)
    covered = 0

    for _ in range(100):          # tối đa 100 patches
        if covered >= target_pixels:
            break
        ph = rng.randint(max(4, H // 8), max(5, H // 3))
        pw = rng.randint(max(4, W // 8), max(5, W // 3))
        py = rng.randint(0, max(1, H - ph))
        px = rng.randint(0, max(1, W - pw))
        arr[py : py + ph, px : px + pw] = 0
        covered += ph * pw

    return Image.fromarray(arr)


def _apply_mask_erosion(img: Image.Image, level: int) -> Image.Image:
    """
    Erode vùng tế bào (non-background) vào trong 'level' pixel từ boundary.
    Mô phỏng: mask boundary bị thu hẹp (under-segmentation).

    level: số pixel erode. Dùng PIL MinFilter(3) mỗi iteration ≈ 1px erosion.
    """
    if level <= 0:
        return img
    arr = np.array(img)
    # Foreground: pixel có giá trị max channel > ngưỡng
    fg = (arr.max(axis=2) > 20).astype(np.uint8) * 255

    mask_pil = Image.fromarray(fg)
    for _ in range(int(level)):
        mask_pil = mask_pil.filter(ImageFilter.MinFilter(3))   # 3×3 min = 1px erosion

    eroded = np.array(mask_pil) > 0
    result = arr.copy()
    result[~eroded] = 0          # Black out các pixel bị erode
    return Image.fromarray(result)


def _apply_background_noise(
    img: Image.Image, level: float, rng: np.random.RandomState
) -> Image.Image:
    """
    Thêm Gaussian noise vào vùng background (pixel ≈ 0).
    Mô phỏng: background lọt vào mask (over-segmentation / imperfect subtraction).

    level: std của noise theo thang 0-255. Ví dụ: 20, 40, 80.
    """
    if level <= 0:
        return img
    arr = np.array(img, dtype=np.float32)
    bg = arr.max(axis=2) < 20    # True = background pixels

    noise = rng.normal(0, level, arr.shape).astype(np.float32)
    result = arr.copy()
    result[bg] = np.clip(result[bg] + noise[bg], 0, 255)
    return Image.fromarray(result.astype(np.uint8))


# =============================================================================
# DegradedTask1Dataset
# =============================================================================

class DegradedTask1Dataset(Dataset):
    """
    Wrapper quanh Task1Dataset — áp dụng degradation on-the-fly trước transform.

    Không load lại paths/labels — tái sử dụng trực tiếp từ base_ds.
    Degradation xảy ra ở không gian PIL (trước Normalize/ToTensor).
    """

    MODES = ('none', 'occlusion', 'erosion', 'bg_noise')

    def __init__(
        self,
        base_ds: Task1Dataset,
        mode: str,
        level: float,
        seed: int = 42,
    ):
        assert mode in self.MODES, f"mode phải là một trong {self.MODES}"
        self.base_ds = base_ds
        self.mode    = mode
        self.level   = level
        self._rng    = np.random.RandomState(seed)

    def __len__(self):
        return len(self.base_ds)

    def __getitem__(self, idx):
        path  = self.base_ds.paths[idx]
        label = self.base_ds.labels[idx]
        orig  = self.base_ds.original_files[idx]
        fname = self.base_ds.filenames[idx]

        try:
            img = Image.open(path).convert('RGB')
        except Exception:
            img = Image.new('RGB', (224, 224), color=0)

        # Apply degradation
        if self.mode == 'occlusion':
            img = _apply_random_occlusion(img, self.level, self._rng)
        elif self.mode == 'erosion':
            img = _apply_mask_erosion(img, int(self.level))
        elif self.mode == 'bg_noise':
            img = _apply_background_noise(img, self.level, self._rng)

        # Apply test transform (Resize → ToTensor → Normalize)
        if self.base_ds.transform:
            img = self.base_ds.transform(img)

        return img, label, orig, fname


class RandomErosionOnTestDataset(Dataset):
    """
    Wrapper quanh Task1Dataset — random erosion cho tập test nhưng KHÔNG áp dụng cho mọi ảnh.

    - Mỗi sample có xác suất p bị erosion.
    - Nếu bị erosion: erosion_px được chọn từ erosion_px_choices (ví dụ 1 hoặc 2).
    - Deterministic theo (base_seed, fname) để kết quả không phụ thuộc num_workers/shuffle.
    """

    def __init__(
        self,
        base_ds: Task1Dataset,
        p: float = 0.5,
        erosion_px_choices: List[int] = None,
        base_seed: int = 0,
    ):
        self.base_ds = base_ds
        self.p = float(p)
        self.erosion_px_choices = erosion_px_choices or [1, 2]
        self.base_seed = int(base_seed)

        if not (0.0 <= self.p <= 1.0):
            raise ValueError("p phải nằm trong [0, 1]")
        if any(int(x) <= 0 for x in self.erosion_px_choices):
            raise ValueError("erosion_px_choices phải là các số nguyên dương")

    def __len__(self):
        return len(self.base_ds)

    def _det_u01(self, fname: str) -> float:
        s = f"{self.base_seed}|{fname}".encode("utf-8", errors="ignore")
        h = hashlib.md5(s).hexdigest()
        # lấy 8 hex chars đầu → 32-bit int
        v = int(h[:8], 16)
        return v / 0xFFFFFFFF

    def _det_choice(self, fname: str, choices: List[int]) -> int:
        s = f"{self.base_seed}|{fname}|choice".encode("utf-8", errors="ignore")
        h = hashlib.md5(s).hexdigest()
        v = int(h[:8], 16)
        return choices[v % len(choices)]

    def __getitem__(self, idx):
        path  = self.base_ds.paths[idx]
        label = self.base_ds.labels[idx]
        orig  = self.base_ds.original_files[idx]
        fname = self.base_ds.filenames[idx]

        try:
            img = Image.open(path).convert('RGB')
        except Exception:
            img = Image.new('RGB', (224, 224), color=0)

        u = self._det_u01(fname)
        if u < self.p:
            px = self._det_choice(fname, self.erosion_px_choices)
            img = _apply_mask_erosion(img, int(px))

        if self.base_ds.transform:
            img = self.base_ds.transform(img)

        return img, label, orig, fname


# =============================================================================
# Default degradation levels
# =============================================================================

DEFAULT_LEVELS = {
    'occlusion': [0.0, 0.10, 0.20, 0.30, 0.40, 0.50],
    'erosion':   [0,   2,    5,    10,   15,   20  ],
    'bg_noise':  [0,   10,   20,   40,   60,   80  ],
}

LEVEL_LABELS = {
    'occlusion': ['0%', '10%', '20%', '30%', '40%', '50%'],
    'erosion':   ['0px', '2px', '5px', '10px', '15px', '20px'],
    'bg_noise':  ['σ=0', 'σ=10', 'σ=20', 'σ=40', 'σ=60', 'σ=80'],
}


# =============================================================================
# Core evaluation: evaluate 1 model × 1 drug × all seeds × all degradation levels
# =============================================================================

def eval_model_robustness(
    drug:        str,
    model_name:  str,
    seeds:       List[int],
    cfg:         Config,
    modes:       List[str],
    device:      torch.device,
    use_bf16:    bool = False,
) -> pd.DataFrame:
    """
    Load checkpoint mỗi seed, test với các degradation levels.

    Returns:
        DataFrame với columns:
            drug, model, mode, level_label, seed,
            image_f1, image_recall, image_precision, image_auc, image_acc,
            cell_f1, cell_recall, cell_auc
    """
    result_dir = cfg.result_dir(1, drug, model_name)
    csv_path   = os.path.join(result_dir, 'run_metrics.csv')

    if not os.path.exists(csv_path):
        print(f"  [WARN] Không tìm thấy run_metrics.csv: {csv_path}")
        return pd.DataFrame()

    # Load optimal threshold per seed từ run_metrics.csv
    metrics_df = pd.read_csv(csv_path)
    img_df     = metrics_df[metrics_df['level'] == 'image_level'].copy()

    # Loại bỏ degenerate threshold (=1.0 hoặc <0.05) trước khi tính median
    # Degenerate xảy ra khi pr_equal tìm được threshold cực đoan → predict all-one-class
    valid_thresh = img_df['opt_threshold']
    valid_thresh = valid_thresh[(valid_thresh > 0.05) & (valid_thresh < 0.99)]
    median_thresh = float(valid_thresh.median()) if len(valid_thresh) > 0 else 0.5

    # Dùng per-seed threshold nếu hợp lệ, ngược lại fallback về median_thresh
    raw_thresholds = img_df.set_index('seed')['opt_threshold'].to_dict()
    thresholds = {
        seed: (t if 0.05 < t < 0.99 else median_thresh)
        for seed, t in raw_thresholds.items()
    }
    print(f"  Thresholds (valid/median={median_thresh:.4f}): { {s: round(t,4) for s,t in thresholds.items()} }")

    # Build base test dataset (split='val' = CSV val → actual test set)
    drug_dir = cfg.drug_dir(drug)
    csv_split = cfg.split_csv(drug)
    full_df  = pd.read_csv(csv_split)
    test_df  = full_df[full_df['split'] == 'val'].reset_index(drop=True)

    transform = get_test_transform(cfg.img_size)
    base_test_ds = Task1Dataset(test_df, drug_dir, split='val', transform=None)
    # transform=None để DegradedTask1Dataset tự apply sau degradation

    # Gán transform vào base_ds để DegradedTask1Dataset có thể dùng
    base_test_ds.transform = transform

    rows = []

    for seed in seeds:
        ckpt_path = cfg.checkpoint_path(1, drug, model_name, seed)
        if not os.path.exists(ckpt_path):
            print(f"  [WARN] Không tìm thấy checkpoint: {ckpt_path} — skip seed {seed}")
            continue

        # Load model
        print(f"  Loading seed={seed} | {model_name} | {drug} ...")
        model = get_model(model_name, num_classes=2, pretrained=False)
        ckpt  = torch.load(ckpt_path, map_location=device)
        model.load_state_dict(ckpt['model_state'])
        model = model.to(device)
        model.eval()

        threshold = thresholds.get(seed, 0.5)

        for mode in modes:
            levels       = DEFAULT_LEVELS[mode]
            level_labels = LEVEL_LABELS[mode]

            for level, level_label in zip(levels, level_labels):
                # Wrap base dataset với degradation
                deg_ds = DegradedTask1Dataset(
                    base_ds = base_test_ds,
                    mode    = mode,
                    level   = level,
                    seed    = seed * 1000 + int(level * 100),  # reproducible per (seed, level)
                )
                loader = DataLoader(
                    deg_ds,
                    batch_size  = cfg.batch_size * 2,
                    num_workers = cfg.num_workers,
                    shuffle     = False,
                    pin_memory  = True,
                )

                t_true, t_pred, t_prob, t_orig, _ = evaluate_loader(
                    model, loader, device, use_bf16=use_bf16
                )

                cell_m = compute_metrics(t_true, t_pred, t_prob, threshold=threshold)
                img_y_true, img_y_pred, img_y_prob, _ = image_level_voting(
                    t_orig, t_true, t_pred, t_prob, method='soft'
                )
                img_m = compute_metrics(img_y_true, img_y_pred, img_y_prob, threshold=threshold)

                rows.append({
                    'drug':              drug,
                    'model':             model_name,
                    'mode':              mode,
                    'level_label':       level_label,
                    'level_value':       level,
                    'seed':              seed,
                    # Image-level
                    'image_f1':          img_m['f1_resistant'],
                    'image_recall':      img_m['recall_resistant'],
                    'image_precision':   img_m['precision_resistant'],
                    'image_auc':         img_m['auc'],
                    'image_accuracy':    img_m['accuracy'],
                    'image_recall_s':    img_m['recall_sensitive'],
                    # Cell-level
                    'cell_f1':           cell_m['f1_resistant'],
                    'cell_recall':       cell_m['recall_resistant'],
                    'cell_auc':          cell_m['auc'],
                })

        # Giải phóng GPU memory
        del model
        torch.cuda.empty_cache()

    return pd.DataFrame(rows)


def eval_model_random_erosion_test(
    drug:        str,
    model_name:  str,
    seeds:       List[int],
    cfg:         Config,
    device:      torch.device,
    p:           float = 0.5,
    px_choices:  List[int] = None,
    use_bf16:    bool = False,
) -> pd.DataFrame:
    """
    Đánh giá 2 điều kiện trên TEST:
      - baseline: không erosion
      - random_erosion: erosion ngẫu nhiên (1–2px) với xác suất p, không áp dụng cho mọi ảnh
    """
    px_choices = px_choices or [1, 2]

    result_dir = cfg.result_dir(1, drug, model_name)
    csv_path   = os.path.join(result_dir, 'run_metrics.csv')
    if not os.path.exists(csv_path):
        print(f"  [WARN] Không tìm thấy run_metrics.csv: {csv_path}")
        return pd.DataFrame()

    metrics_df = pd.read_csv(csv_path)
    img_df     = metrics_df[metrics_df['level'] == 'image_level'].copy()
    valid_thresh = img_df['opt_threshold']
    valid_thresh = valid_thresh[(valid_thresh > 0.05) & (valid_thresh < 0.99)]
    median_thresh = float(valid_thresh.median()) if len(valid_thresh) > 0 else 0.5
    raw_thresholds = img_df.set_index('seed')['opt_threshold'].to_dict()
    thresholds = {
        seed: (t if 0.05 < t < 0.99 else median_thresh)
        for seed, t in raw_thresholds.items()
    }
    print(f"  Thresholds (valid/median={median_thresh:.4f}): { {s: round(t,4) for s,t in thresholds.items()} }")

    drug_dir = cfg.drug_dir(drug)
    csv_split = cfg.split_csv(drug)
    full_df  = pd.read_csv(csv_split)
    test_df  = full_df[full_df['split'] == 'val'].reset_index(drop=True)

    transform = get_test_transform(cfg.img_size)
    base_test_ds = Task1Dataset(test_df, drug_dir, split='val', transform=None)
    base_test_ds.transform = transform

    rows = []
    for seed in seeds:
        ckpt_path = cfg.checkpoint_path(1, drug, model_name, seed)
        if not os.path.exists(ckpt_path):
            print(f"  [WARN] Không tìm thấy checkpoint: {ckpt_path} — skip seed {seed}")
            continue

        print(f"  Loading seed={seed} | {model_name} | {drug} ...")
        model = get_model(model_name, num_classes=2, pretrained=False)
        ckpt  = torch.load(ckpt_path, map_location=device)
        model.load_state_dict(ckpt['model_state'])
        model = model.to(device)
        model.eval()

        threshold = thresholds.get(seed, 0.5)

        conditions = [
            ('baseline', base_test_ds),
            ('rand_erosion', RandomErosionOnTestDataset(
                base_ds=base_test_ds,
                p=p,
                erosion_px_choices=[int(x) for x in px_choices],
                base_seed=seed,
            )),
        ]

        for cond_name, ds in conditions:
            loader = DataLoader(
                ds,
                batch_size  = cfg.batch_size * 2,
                num_workers = cfg.num_workers,
                shuffle     = False,
                pin_memory  = True,
            )
            t_true, t_pred, t_prob, t_orig, _ = evaluate_loader(
                model, loader, device, use_bf16=use_bf16
            )
            cell_m = compute_metrics(t_true, t_pred, t_prob, threshold=threshold)
            img_y_true, img_y_pred, img_y_prob, _ = image_level_voting(
                t_orig, t_true, t_pred, t_prob, method='soft'
            )
            img_m = compute_metrics(img_y_true, img_y_pred, img_y_prob, threshold=threshold)

            rows.append({
                'drug':              drug,
                'model':             model_name,
                'mode':              f"random_erosion(p={p},px={','.join(map(str,px_choices))})",
                'level_label':       cond_name,
                'level_value':       0 if cond_name == 'baseline' else 1,
                'seed':              seed,
                'image_f1':          img_m['f1_resistant'],
                'image_recall':      img_m['recall_resistant'],
                'image_precision':   img_m['precision_resistant'],
                'image_auc':         img_m['auc'],
                'image_accuracy':    img_m['accuracy'],
                'image_recall_s':    img_m['recall_sensitive'],
                'cell_f1':           cell_m['f1_resistant'],
                'cell_recall':       cell_m['recall_resistant'],
                'cell_auc':          cell_m['auc'],
            })

        del model
        torch.cuda.empty_cache()

    return pd.DataFrame(rows)


# =============================================================================
# Aggregate: mean ± std across seeds per (mode, level)
# =============================================================================

def aggregate_results(df: pd.DataFrame) -> pd.DataFrame:
    """Tính mean ± std across seeds, per (drug, model, mode, level)."""
    if df.empty:
        return df

    metric_cols = [
        'image_f1', 'image_recall', 'image_precision',
        'image_auc', 'image_accuracy', 'image_recall_s',
        'cell_f1', 'cell_recall', 'cell_auc',
    ]
    group_cols = ['drug', 'model', 'mode', 'level_label', 'level_value']

    agg = (
        df.groupby(group_cols)[metric_cols]
        .agg(['mean', 'std'])
        .reset_index()
    )
    # Flatten multi-level columns
    agg.columns = [
        '_'.join(c).rstrip('_') if isinstance(c, tuple) else c
        for c in agg.columns
    ]
    # Tạo cột "mean ± std" cho từng metric
    for col in metric_cols:
        m = f'{col}_mean'
        s = f'{col}_std'
        if m in agg.columns and s in agg.columns:
            agg[f'{col}_str'] = agg.apply(
                lambda r: f"{r[m]:.4f} ± {r[s]:.4f}", axis=1
            )
    return agg


# =============================================================================
# Print table
# =============================================================================

def print_robustness_table(agg: pd.DataFrame, mode: str, drug: str, model: str):
    """In bảng degradation level → metrics đẹp ra terminal."""
    sub = agg[(agg['mode'] == mode) & (agg['drug'] == drug) & (agg['model'] == model)]
    if sub.empty:
        return

    label_map = {'occlusion': 'Occlusion %', 'erosion': 'Erosion (px)', 'bg_noise': 'BG Noise σ'}
    mode_label = label_map.get(mode, mode)

    print(f"\n{'═'*100}")
    print(f"  [{drug}] [{model}] — {mode_label} Robustness")
    print(f"{'═'*100}")
    print(f"  {'Level':<10}  {'Image F1':>10}  {'Image Recall':>14}  {'Image AUC':>12}  {'Cell F1':>10}")
    print(f"  {'-'*96}")
    for _, row in sub.iterrows():
        f1  = row.get('image_f1_str', '')
        rec = row.get('image_recall_str', '')
        auc = row.get('image_auc_str', '')
        cf1 = row.get('cell_f1_str', '')
        print(f"  {row['level_label']:<10}  {f1:>22}  {rec:>22}  {auc:>20}  {cf1:>22}")
    print(f"{'═'*100}\n")


# =============================================================================
# Save Excel
# =============================================================================

def save_excel_robustness(agg: pd.DataFrame, out_path: str):
    """Lưu kết quả robustness ra Excel với format đẹp, mỗi mode = 1 sheet."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("  [Warning] openpyxl chưa cài. Lưu CSV thay thế.")
        agg.to_csv(out_path.replace('.xlsx', '.csv'), index=False)
        print(f"  ✓ Saved CSV → {out_path.replace('.xlsx', '.csv')}")
        return

    HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
    HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    BODY_FONT   = Font(name="Calibri", size=10)
    BEST_FILL   = PatternFill("solid", fgColor="E2EFDA")   # Xanh lá (tốt nhất)
    WORST_FILL  = PatternFill("solid", fgColor="FDECEA")   # Đỏ nhạt (tệ nhất)
    ALT_FILL    = [
        PatternFill("solid", fgColor="D6E4F0"),
        PatternFill("solid", fgColor="FFFFFF"),
    ]
    CENTER = Alignment(horizontal="center", vertical="center")
    LEFT   = Alignment(horizontal="left",   vertical="center")
    BORDER = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )

    SHOW_METRICS = [
        ('Image F1 (R)',        'image_f1_str',        'image_f1_mean'),
        ('Image Recall (R)',    'image_recall_str',    'image_recall_mean'),
        ('Image Precision (R)','image_precision_str', 'image_precision_mean'),
        ('Image AUC',          'image_auc_str',       'image_auc_mean'),
        ('Image Accuracy',     'image_accuracy_str',  'image_accuracy_mean'),
        ('Image Spec (S)',      'image_recall_s_str',  'image_recall_s_mean'),
        ('Cell F1 (R)',         'cell_f1_str',         'cell_f1_mean'),
        ('Cell Recall (R)',     'cell_recall_str',     'cell_recall_mean'),
        ('Cell AUC',            'cell_auc_str',        'cell_auc_mean'),
    ]

    wb = openpyxl.Workbook()
    first = True

    drugs   = agg['drug'].unique()
    models  = agg['model'].unique()
    modes   = agg['mode'].unique()

    for mode in modes:
        for drug in drugs:
            for model_name in models:
                sub = agg[
                    (agg['mode'] == mode) &
                    (agg['drug'] == drug) &
                    (agg['model'] == model_name)
                ].sort_values('level_value')
                if sub.empty:
                    continue

                sheet_title = f"{mode[:4]}-{drug}-{model_name[:8]}"[:31]
                if first:
                    ws = wb.active
                    ws.title = sheet_title
                    first = False
                else:
                    ws = wb.create_sheet(sheet_title)

                # Title
                title = f"Robustness ({mode}) — {drug} — {model_name}"
                ws.append([title])
                ws.merge_cells(start_row=1, start_column=1,
                               end_row=1, end_column=2 + len(SHOW_METRICS))
                tc = ws.cell(1, 1)
                tc.font      = Font(name="Calibri", bold=True, size=12, color="1F4E79")
                tc.alignment = CENTER
                ws.row_dimensions[1].height = 18

                # Header
                hrow_data = ['Level'] + [m[0] for m in SHOW_METRICS]
                ws.append(hrow_data)
                hrow = ws.max_row
                for ci, v in enumerate(hrow_data, 1):
                    c = ws.cell(hrow, ci, v)
                    c.font      = HEADER_FONT
                    c.fill      = HEADER_FILL
                    c.alignment = CENTER
                    c.border    = BORDER
                ws.row_dimensions[hrow].height = 26

                # Data rows
                baseline_f1 = None
                data_rows = []
                for ri, (_, row) in enumerate(sub.iterrows()):
                    data_row = [row['level_label']] + [
                        row.get(m[1], '') for m in SHOW_METRICS
                    ]
                    data_rows.append((ri, row, data_row))
                    if ri == 0:
                        baseline_f1 = row.get('image_f1_mean', None)

                for ri, row, data_row in data_rows:
                    ws.append(data_row)
                    drow = ws.max_row
                    base_fill = ALT_FILL[ri % 2]
                    for ci, val in enumerate(data_row, 1):
                        c = ws.cell(drow, ci, val)
                        c.font      = BODY_FONT
                        c.fill      = base_fill
                        c.alignment = LEFT if ci == 1 else CENTER
                        c.border    = BORDER
                    # Highlight: baseline=BEST, worst row=WORST
                    if ri == 0 and len(data_rows) > 1:
                        for ci in range(1, len(data_row) + 1):
                            ws.cell(drow, ci).fill = BEST_FILL
                    elif ri == len(data_rows) - 1 and len(data_rows) > 1:
                        for ci in range(1, len(data_row) + 1):
                            ws.cell(drow, ci).fill = WORST_FILL

                # Column widths
                widths = [10] + [20] * len(SHOW_METRICS)
                for i, w in enumerate(widths, 1):
                    ws.column_dimensions[get_column_letter(i)].width = w
                ws.freeze_panes = ws.cell(hrow + 1, 1)

    # Sheet tổng hợp: tất cả mode × drug × model cạnh nhau
    ws_summary = wb.create_sheet("Summary")
    _write_summary_sheet(ws_summary, agg, SHOW_METRICS,
                         HEADER_FILL, HEADER_FONT, BODY_FONT, CENTER, LEFT, BORDER,
                         BEST_FILL, WORST_FILL, ALT_FILL)

    wb.save(out_path)
    print(f"✓ Đã lưu Excel → {out_path}")


def _write_summary_sheet(ws, agg, show_metrics,
                          header_fill, header_font, body_font,
                          center, left, border, best_fill, worst_fill, alt_fill):
    """Sheet tổng hợp: mỗi hàng = (drug, model, mode, level), cols = metrics."""
    from openpyxl.styles import Font, PatternFill
    header = ['Drug', 'Model', 'Mode', 'Level'] + [m[0] for m in show_metrics]
    ws.append(header)
    hrow = ws.max_row
    for ci, v in enumerate(header, 1):
        c = ws.cell(hrow, ci, v)
        c.font      = header_font
        c.fill      = header_fill
        c.alignment = center
        c.border    = border
    ws.row_dimensions[hrow].height = 26

    prev_group = None
    grp_idx = 0
    for _, row in agg.sort_values(['drug', 'model', 'mode', 'level_value']).iterrows():
        group = (row['drug'], row['model'], row['mode'])
        if group != prev_group:
            prev_group = group
            grp_idx += 1
        data = [
            row['drug'], row['model'], row['mode'], row['level_label'],
        ] + [row.get(m[1], '') for m in show_metrics]
        ws.append(data)
        drow = ws.max_row
        fill = alt_fill[grp_idx % 2]
        for ci, v in enumerate(data, 1):
            c = ws.cell(drow, ci, v)
            c.font      = body_font
            c.fill      = fill
            c.alignment = left if ci <= 4 else center
            c.border    = border

    from openpyxl.utils import get_column_letter
    widths = [8, 22, 10, 10] + [20] * len(show_metrics)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(hrow + 1, 1)


# =============================================================================
# Main
# =============================================================================

def parse_args():
    parser = argparse.ArgumentParser(
        description='Đánh giá robustness của Task1 model với degraded mask'
    )
    parser.add_argument(
        '--drug', type=str, nargs='+', default=['CIP'],
        help='Drug(s) cần đánh giá. Default: CIP',
    )
    parser.add_argument(
        '--model', type=str, nargs='+',
        default=['efficientnet_b5'],
        help='Model name(s) (phải trùng với thư mục results/task1/{drug}/{model}). Default: efficientnet_b5',
    )
    parser.add_argument(
        '--seeds', type=int, nargs='+', default=[0, 1, 2, 3, 4],
        help='Seed(s) cần eval. Default: 0 1 2 3 4',
    )
    parser.add_argument(
        '--mode', type=str, nargs='+',
        choices=['occlusion', 'erosion', 'bg_noise', 'all'],
        default=['all'],
        help='Degradation mode(s). Default: all',
    )
    parser.add_argument(
        '--random_erosion_test',
        action='store_true', default=False,
        help='[Task1][TEST] Random erosion 1-2px với xác suất p trên tập test (không áp dụng cho mọi ảnh). '
             'Nếu bật, sẽ chỉ chạy 2 điều kiện: baseline vs rand_erosion.',
    )
    parser.add_argument(
        '--random_erosion_px', type=int, nargs='+', default=[1, 2],
        help='[Task1][TEST] Các mức erosion px dùng cho random_erosion_test. Default: 1 2',
    )
    parser.add_argument(
        '--random_erosion_p', type=float, default=0.5,
        help='[Task1][TEST] Xác suất mỗi ảnh bị erosion khi random_erosion_test. Default: 0.5',
    )
    parser.add_argument(
        '--out', type=str, default=None,
        help='Đường dẫn file output Excel. Default: degraded_mask_{drug}_{model}.xlsx',
    )
    parser.add_argument(
        '--results_dir', type=str, default='/home/aisia/new_mia/results',
        help='Thư mục chứa kết quả training (để tìm checkpoint + run_metrics.csv)',
    )
    parser.add_argument(
        '--data_root', type=str, default='/home/aisia/new_mia',
        help='Root thư mục chứa dữ liệu (*_joblib)',
    )
    parser.add_argument(
        '--batch_size', type=int, default=128,
        help='Batch size cho inference. Default: 128',
    )
    parser.add_argument(
        '--num_workers', type=int, default=8,
        help='Số worker cho DataLoader. Default: 8',
    )
    parser.add_argument(
        '--bf16', action='store_true', default=False,
        help='Dùng BFloat16 inference (RTX 30xx/40xx)',
    )
    return parser.parse_args()


def main():
    args = parse_args()

    # Expand 'all' → tất cả modes
    if 'all' in args.mode:
        modes = ['occlusion', 'erosion', 'bg_noise']
    else:
        modes = args.mode

    # Config
    cfg = Config()
    cfg.data_root   = args.data_root
    cfg.output_dir  = args.results_dir
    cfg.batch_size  = args.batch_size
    cfg.num_workers = args.num_workers
    cfg.use_bf16    = args.bf16

    device = torch.device('cuda' if torch.cuda.is_available() else 'cpu')
    print(f"\nDevice: {device} | BF16: {args.bf16}")
    print(f"Drugs:  {args.drug}")
    print(f"Models: {args.model}")
    print(f"Seeds:  {args.seeds}")
    print(f"Modes:  {modes}\n")

    all_raw_dfs = []

    for drug in args.drug:
        for model_name in args.model:
            print(f"\n{'═'*60}")
            print(f"  Evaluating: {drug} | {model_name}")
            print(f"{'═'*60}")

            if args.random_erosion_test:
                raw_df = eval_model_random_erosion_test(
                    drug       = drug,
                    model_name = model_name,
                    seeds      = args.seeds,
                    cfg        = cfg,
                    device     = device,
                    p          = args.random_erosion_p,
                    px_choices = args.random_erosion_px,
                    use_bf16   = args.bf16,
                )
            else:
                raw_df = eval_model_robustness(
                    drug       = drug,
                    model_name = model_name,
                    seeds      = args.seeds,
                    cfg        = cfg,
                    modes      = modes,
                    device     = device,
                    use_bf16   = args.bf16,
                )
            if not raw_df.empty:
                all_raw_dfs.append(raw_df)

    if not all_raw_dfs:
        print("Không có kết quả. Kiểm tra lại đường dẫn checkpoints và data.")
        return

    full_raw = pd.concat(all_raw_dfs, ignore_index=True)
    agg      = aggregate_results(full_raw)

    # In bảng ra terminal
    for drug in args.drug:
        for model_name in args.model:
            if args.random_erosion_test:
                # Mode string có dạng random_erosion(...)
                for mode in agg['mode'].unique():
                    print_robustness_table(agg, mode=mode, drug=drug, model=model_name)
            else:
                for mode in modes:
                    print_robustness_table(agg, mode=mode, drug=drug, model=model_name)

    # Lưu Excel
    tag = f"{'_'.join(args.drug)}_{'_'.join(m[:3] for m in args.model)}"
    out_path = args.out or f"degraded_mask_robustness_{tag}.xlsx"
    save_excel_robustness(agg, out_path)

    # Lưu thêm CSV raw (per seed)
    csv_path = out_path.replace('.xlsx', '_raw.csv')
    full_raw.to_csv(csv_path, index=False)
    print(f"✓ Đã lưu CSV raw → {csv_path}")


if __name__ == '__main__':
    main()
